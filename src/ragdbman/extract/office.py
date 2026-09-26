# SPDX-FileCopyrightText: 2026 Bela Istvan MIHALIK
# SPDX-License-Identifier: Apache-2.0

"""OOXML, ODF, EPUB and Kindle adapters with safe XML and ZIP reads."""

import posixpath
import re
import shutil
from pathlib import Path
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring

from ..errors import RagError
from ..models import Block, Document
from .text import html

MAX_UNPACKED = 512 * 1024 * 1024


def checked_zip(path: Path) -> ZipFile:
    archive = ZipFile(path)
    if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED:
        archive.close()
        raise RagError("FILE_TOO_LARGE", "Archive's uncompressed contents exceed 512 MiB")
    return archive


def local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def text_nodes(node, tags=("t",)):
    return "".join(child.text or "" for child in node.iter() if local(child.tag) in tags)


def docx(path: Path) -> Document:
    with checked_zip(path) as z:
        root = fromstring(z.read("word/document.xml"))
    blocks, section = [], None
    for paragraph in root.iter():
        if local(paragraph.tag) != "p":
            continue
        content = text_nodes(paragraph)
        if not content.strip():
            continue
        style = next(
            (next(iter(n.attrib.values()), "") for n in paragraph.iter() if local(n.tag) == "pStyle"), ""
        )
        heading = style.lower().startswith("heading")
        if heading:
            section = content
        blocks.append(Block(content, is_heading=heading, heading=section, section_path=section))
    return Document(blocks, "docx", title=path.stem)


def pptx(path: Path) -> Document:
    blocks = []
    with checked_zip(path) as z:
        slides = sorted(
            (n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"slide(\d+)", n)[1]),
        )
        for n, member in enumerate(slides, 1):
            root = fromstring(z.read(member))
            for para in root.iter():
                if local(para.tag) == "p":
                    text = text_nodes(para)
                    if text.strip():
                        blocks.append(Block(text, slide=n, section_path=f"Slide {n}", atomic=True))
    return Document(blocks, "pptx", title=path.stem)


def xlsx(path: Path) -> Document:
    from openpyxl import load_workbook

    with checked_zip(path):
        pass
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    try:
        for sheet in wb:
            for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                text = "\t".join("" if cell is None else str(cell) for cell in row)
                if text.strip():
                    blocks.append(
                        Block(
                            text,
                            sheet=sheet.title,
                            section_path=sheet.title,
                            line_start=n,
                            line_end=n,
                            atomic=True,
                        )
                    )
    finally:
        wb.close()
    return Document(blocks, "xlsx", title=path.stem)


def odf(path: Path) -> Document:
    with checked_zip(path) as z:
        root = fromstring(z.read("content.xml"))
    blocks = []
    ext = path.suffix.lower()
    if ext == ".ods":
        for table in root.iter():
            if local(table.tag) != "table":
                continue
            name = next((v for k, v in table.attrib.items() if local(k) == "name"), "Sheet")
            rownum = 0
            for row in table:
                if local(row.tag) != "table-row":
                    continue
                repeats = int(
                    next((v for k, v in row.attrib.items() if local(k) == "number-rows-repeated"), "1")
                )
                values = []
                for cell in row:
                    value = " ".join("".join(p.itertext()) for p in cell if local(p.tag) == "p")
                    if not value:
                        value = next((v for k, v in cell.attrib.items() if local(k) == "value"), "")
                    count = int(
                        next(
                            (v for k, v in cell.attrib.items() if local(k) == "number-columns-repeated"), "1"
                        )
                    )
                    if count > 10000:
                        if value:
                            raise RagError("FILE_TOO_LARGE", "ODF column repetition exceeds safety limit")
                        count = 1
                    values.extend([value] * count)
                text = "\t".join(values).rstrip()
                if text.strip():
                    if repeats > 10000:
                        raise RagError("FILE_TOO_LARGE", "ODF row repetition exceeds safety limit")
                    for i in range(repeats):
                        blocks.append(
                            Block(
                                text,
                                sheet=name,
                                section_path=name,
                                atomic=True,
                                line_start=rownum + i + 1,
                                line_end=rownum + i + 1,
                            )
                        )
                rownum += repeats
    elif ext == ".odp":
        for num, page in enumerate((p for p in root.iter() if local(p.tag) == "page"), 1):
            for node in page.iter():
                if local(node.tag) in {"p", "h"}:
                    text = "".join(node.itertext())
                    if text.strip():
                        blocks.append(Block(text, slide=num, section_path=f"Slide {num}", atomic=True))
    else:
        section = None
        for node in root.iter():
            if local(node.tag) in {"p", "h"}:
                text = "".join(node.itertext())
                heading = local(node.tag) == "h"
                if heading:
                    section = text
                if text.strip():
                    blocks.append(Block(text, is_heading=heading, heading=section, section_path=section))
    return Document(blocks, ext.lstrip("."), title=path.stem)


def epub(path: Path) -> Document:
    blocks = []
    with checked_zip(path) as z:
        container = fromstring(z.read("META-INF/container.xml"))
        opf_name = next(n.attrib["full-path"] for n in container.iter() if local(n.tag) == "rootfile")
        opf = fromstring(z.read(opf_name))
        manifest = {n.attrib["id"]: n.attrib["href"] for n in opf.iter() if local(n.tag) == "item"}
        spine = [manifest[n.attrib["idref"]] for n in opf.iter() if local(n.tag) == "itemref"]
        title = next((n.text for n in opf.iter() if local(n.tag) == "title"), path.stem)
        for member in spine:
            name = posixpath.normpath(posixpath.join(posixpath.dirname(opf_name), member.split("#")[0]))
            document = html(z.read(name).decode("utf-8", errors="replace"))
            for block in document.blocks:
                block.section_path = block.section_path or document.title or member
            blocks.extend(document.blocks)
    return Document(blocks, "epub", title=title)


def azw3(path: Path) -> Document:
    import mobi

    directory = None
    try:
        directory, output = mobi.extract(str(path))
        output = Path(output)
        doc = epub(output) if output.suffix.lower() == ".epub" else html(output.read_text(errors="replace"))
        doc.extractor_name = "azw3"
        return doc
    except Exception as exc:
        raise RagError("EXTRACTION_FAILED", f"Kindle extraction failed (DRM is unsupported): {exc}") from exc
    finally:
        if directory:
            shutil.rmtree(directory, ignore_errors=True)
